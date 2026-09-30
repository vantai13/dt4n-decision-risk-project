"""P1v2/L1.10 — thêm sheet 'Novelty Matrix v4' vào novelty_matrix.xlsx.

Idempotent: chạy lại sẽ ghi đè sheet v4, không đụng các sheet cũ (v3 là lịch sử).
Cột mới theo PHASE_1v2 L1.10 việc 4. Provenance: Claude (AI) soạn 2026-09-30; tác giả kiểm từng ô.
Ô chưa kiểm ghi rõ 'Chưa…' — không điền cho đẹp.
"""
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WB = Path(__file__).resolve().parents[1] / "novelty_matrix.xlsx"
SHEET = "Novelty Matrix v4"
NA = "Không áp dụng"

HEADERS = {
    "paper": "Paper / Work",
    "role": "Họ / vai trò",
    "vs_tuned": "Đo giá trị bất định so với ngưỡng đã tune?",
    "oracle": "Oracle cùng thông tin?",
    "mechanism": "Cơ chế tạo bất định trực giao?",
    "decomp": "Phân rã thông tin / an toàn / thích nghi?",
    "diff": "Khác biệt problem/method so với đề tài (D18)",
    "no_claim": "Không được claim",
    "fulltext": "Đã kiểm full text (ngày, §)",
}
KEYS = list(HEADERS)
WIDTHS = [36, 24, 42, 28, 38, 32, 44, 40, 30]

ROWS = [
    dict(paper="DT4N v2 — đề tài (khung Phase 1 v2)", role="Đề tài",
         vs_tuned="Có — K2(α) so với S0 / S0dir / SC trên cùng F, cùng α (F7 §2.3)",
         oracle="Có — oracle bin cùng F dọc quỹ đạo tham chiếu (K16); KHÔNG luôn là cận trên (F8: T10/FH −3,1 SE)",
         mechanism="Sơ bộ — mức tải × độ cong (f05c ±26%); dị loại bị ngưỡng theo chiều hấp thụ (F7); "
                   "lịch sử không tăng κ trong DES (F8, khám phá)",
         decomp="Có — thông tin 59–84% · an toàn 6–14% · tâm 1–4% · thuần ≤ 1,5% (F7 §3–4)",
         diff="—", no_claim="Nguyên lý đổi thứ tự và luật bậc hai (xem các dòng cổ điển)", fulltext=NA),
    dict(paper="Jewson 2003 — Moment based methods for ensemble assessment and calibration (arXiv:physics/0309042)",
         role="HỌ MỚI: dự báo tổ hợp — giá trị của spread",
         vs_tuned="Gần nhất về khái niệm: chỉ-tâm (độ rộng hằng) vs tâm + spread, ngoài mẫu (§3, §4.4–4.5); "
                  "thước đo likelihood, KHÔNG phải giá trị quyết định",
         oracle="Không", mechanism="Một phần: COVS sau hiệu chỉnh nhỏ (§4.2); không cơ chế vật lý", decomp="Không",
         diff="Problem: quyết định nhị phân có ngân sách harm (β ∝ λ); Method: oracle cùng thông tin + phân rã "
              "(đã kiểm full text)",
         no_claim="'Độ rộng thêm ít khi biến thiên dự đoán được của bất định nhỏ; giá trị nằm ở tâm'",
         fulltext="2026-09-30; §1–5 (hình: chú thích)"),
    dict(paper="Houtekamer 1993 (MWR 121(6)); Whitaker & Loughe 1998 (MWR 126)", role="Dự báo tổ hợp — spread–skill",
         vs_tuned="Không trực tiếp: đo tương quan spread–skill", oracle="Không",
         mechanism="Khái niệm: tương quan cực đại khi spread biến thiên mạnh (spread log-normal) ~ κ", decomp="Không",
         diff="Chưa kiểm full text",
         no_claim="'Giá trị bất định phụ thuộc độ phân tán của log s' là ý mới",
         fulltext="Chưa; abstract + thứ cấp 2026-09-30"),
    dict(paper="Birge 1982 — Value of the stochastic solution (Math. Prog. 24(1):314–325)",
         role="Quy hoạch ngẫu nhiên — EVPI, VSS",
         vs_tuned="Khái niệm: VSS = dùng phân phối vs cắm kỳ vọng (KHÔNG tune)",
         oracle="WS = thông tin hoàn hảo (không phải cùng thông tin)", mechanism="Không",
         decomp="Hai bậc: EVPI (thông tin) + VSS (dùng phân phối)",
         diff="Method: baseline đã tune; thêm bậc an toàn và tâm; oracle cùng thông tin — CHƯA kiểm full text Birge",
         no_claim="Ý tưởng tách 'giá trị thông tin hoàn hảo' và 'giá trị dùng phân phối'",
         fulltext="Chưa; định nghĩa qua ≥ 3 nguồn thứ cấp 2026-09-30"),
    dict(paper="von Winterfeldt & Edwards 1973 — Flat maxima (TR 011313-4-T, U. Michigan)",
         role="Toán cổ điển — flat maximum", vs_tuned=NA, oracle=NA, mechanism=NA, decomp=NA,
         diff=NA + " (nền lý thuyết)", no_claim="'Lệch gần tối ưu rẻ; tổn thất bậc hai'",
         fulltext="Chưa; thư mục 2026-09-30"),
    dict(paper="Radner & Stiglitz 1984; Chade & Schlee 2002 (JET 107(2):421–452)",
         role="Toán cổ điển — giá trị biên thông tin = 0 tại 0", vs_tuned=NA, oracle=NA, mechanism=NA, decomp=NA,
         diff=NA + " (nền lý thuyết)", no_claim="'Giá trị thông tin nhỏ là bậc hai'",
         fulltext="Chưa; thư mục + abstract 2026-09-30"),
    dict(paper="Mammen & Tsybakov 1999; Audibert & Tsybakov 2007 (AoS 35(2):608–633)",
         role="Toán cổ điển — điều kiện margin", vs_tuned=NA, oracle=NA, mechanism=NA, decomp=NA,
         diff=NA + " (nền lý thuyết)", no_claim="'Tổn thất ∝ mật độ gần biên × sai lệch²'",
         fulltext="Chưa; abstract 2026-09-30"),
    dict(paper="Tiwari, Kar & Tiwari 2026 — DT belief-state RL, latency-robust ISAC (arXiv:2604.25967v1)",
         role="DT dựng belief state từ telemetry trễ",
         vs_tuned="Không; belief = ước lượng EKF (eq. 13)", oracle="Không", mechanism="Không",
         decomp="Không; [I] DT-only 86% vs 88% ở 100 ms (§IV-D) ⇒ lợi ích chủ yếu ở tâm",
         diff="Problem: hành động liên tục, không đổi/giữ, không ngân sách harm; Method: không oracle (đã kiểm)",
         no_claim="'Twin dựng trạng thái hiện tại từ telemetry trễ để quyết định'",
         fulltext="2026-09-30; §I–V (hình: chú thích)"),
    dict(paper="Li, Wu, Lee & Sun 2026 — From Freshness to Effectiveness (arXiv:2504.19507v3)",
         role="AoI/VoI cho quyết định từ xa",
         vs_tuned="Không: so co-design với các chính sách lấy mẫu chuẩn (§VIII-A)", oracle="Không (MDP biết mô hình)",
         mechanism="Không; tuổi là thông tin phụ trong thống kê đủ (Lemma 1)",
         decomp="Không; ngưỡng tần suất lấy mẫu (Thm 9, Cor. 2)",
         diff="Problem: nguồn Markov hữu hạn, đồng thời chọn lấy mẫu; Method: không ngưỡng đã tune, không oracle "
              "cùng thông tin (đã kiểm)",
         no_claim="'Tuổi là thông tin phụ cho quyết định'; 'độ tươi ≠ giá trị quyết định'",
         fulltext="2026-09-30; §I–VII; §VIII-A/B (hình: chú thích)"),
    dict(paper="Scherrer, Perrig & Schmid 2020 — Value of Information in Selfish Routing (SIROCCO)",
         role="VoI trong định tuyến — tập thể (Wardrop)", vs_tuned="Không", oracle="Không", mechanism="Không",
         decomp="Không", diff="Chế độ khác (tải nội sinh) — chưa kiểm full text",
         no_claim="'Nhiều thông tin hơn có thể không giúp khi nhiều luồng cùng chọn đường'",
         fulltext="Chưa; abstract 2026-09-30"),
    dict(paper="Keralapura, Chuah, Taft & Iannaccone 2008 — Race conditions in coexisting overlays (ToN 16(1))",
         role="Dao động nhiều overlay — tập thể", vs_tuned="Không", oracle="Không", mechanism="Không", decomp="Không",
         diff="Chế độ khác (nhiều overlay) — chưa kiểm full text",
         no_claim="'Nhiều bộ chọn đường độc lập có thể đồng bộ và dao động'",
         fulltext="Chưa; abstract/đoạn trích 2026-09-30"),
    dict(paper="Seshadri & Katz 2003 — Dynamics of simultaneous overlay network routing (TR; WIRED)",
         role="Hysteresis H, overlay tập thể",
         vs_tuned="Không trực tiếp: quét H cố định, MIMD-H, ngẫu nhiên hoá (§IV–V)", oracle="Không",
         mechanism="Gián tiếp: H tối ưu phụ thuộc tham số (Fig. 4–5)", decomp="Không",
         diff="Problem: tập thể, metric mất gói; Method: không oracle (đã kiểm 27/09)",
         no_claim="'H tối ưu phụ thuộc tham số'; 'H thích nghi (MIMD)'", fulltext="2026-09-27; §I–VII"),
    dict(paper="Fischer & Vöcking 2005/2009 — Adaptive routing with stale information",
         role="Thông tin cũ gây dao động (tập thể)", vs_tuned="Không", oracle="Không", mechanism="Không",
         decomp="Không", diff="Problem: tải nội sinh, fluid limit (đã kiểm 27/09)",
         no_claim="'Thông tin cũ làm luật tham lam dao động'", fulltext="2026-09-27; TR §1–4"),
    dict(paper="Liyanage et al. 2026 — Risk-aware stable server selection (arXiv:2604.21483v1)",
         role="(μ, σ) + hysteresis",
         vs_tuned="Không tách được: thiếu ô 'trung bình + hysteresis' (Table II); một trace, không CI",
         oracle="Không", mechanism="Không", decomp="Không",
         diff="Problem: SLO, không ngân sách harm; Method: không oracle, không factorial (đã kiểm 27/09)",
         no_claim="'Dùng bất định + hysteresis để chọn đích'", fulltext="2026-09-27; §I–VI"),
    dict(paper="Burbano et al. 2025 — MO-HAN (arXiv:2511.10146v1)", role="Dự đoán + hysteresis",
         vs_tuned="Không: hysteresis trên điểm có hiệu chỉnh (Eq. 4–7)", oracle="Không", mechanism="Không",
         decomp="Không", diff="Problem: không tuổi/nhiễu; Method: không oracle, không CI (đã kiểm 27/09)",
         no_claim="'Dự đoán + hysteresis'", fulltext="2026-09-27; §I–V"),
    dict(paper="Almohammedi et al. 2026 — Twin-fidelity-aware xApp arbitration (arXiv:2607.22857v1)",
         role="Fallback theo fidelity twin", vs_tuned="Không: fallback khi EWMA sai số ≥ τ (Eq. 12)",
         oracle="Không", mechanism="Không (drift toàn cục)", decomp="Không",
         diff="Problem: lệch mô hình toàn cục (đã kiểm 27/09)", no_claim="'Fallback theo fidelity twin'",
         fulltext="2026-09-27; §I–VIII"),
    dict(paper="OpenTwin v2 2026 (arXiv:2605.24662v2)", role="Cổng an toàn hành động cho DT",
         vs_tuned="Không: gate an toàn tuyệt đối C_t(a) ⊆ S (§VI)", oracle="Không",
         mechanism="Không (drift qua hiệu chỉnh trực tuyến)", decomp="Không",
         diff="Problem: an toàn một hành động (đã kiểm 23/09)", no_claim="'Cổng tin cậy/conformal cho DT'",
         fulltext="2026-09-23; §V–VIII"),
    dict(paper="CERT 2026; LEC 2026; Zhu et al. 2026; Lekeufack et al. 2024 (chi tiết: sheet v3)",
         role="Gate/certificate cho quyết định (RQ2 cũ)",
         vs_tuned="Không: hiệu chỉnh rủi ro của gate", oracle="Không", mechanism="Không", decomp="Không",
         diff="Trùng một phần gate/luật — lý do DP1 = NARROW (F6)", no_claim="Xem v3", fulltext="Xem v3 (23/09)"),
    dict(paper="Veeravalli & Kelly 1997 — A locally optimal handoff algorithm (IEEE TVT 46(3):603–609)",
         role="HỌ HANDOVER — luật so xác suất vs hysteresis đã tune",
         vs_tuned="CÓ, gần nhất: LO vs hysteresis và hysteresis-threshold (2 tham số) trên CÙNG bộ dự báo, cùng số "
                  "handoff; LO ≈ hysteresis-threshold tốt nhất (§V, Fig. 3)",
         oracle="Không cùng thông tin: DP tối ưu biết trước quỹ đạo, dùng làm mốc (§III, Fig. 3)",
         mechanism="Không: phương sai có điều kiện không phụ thuộc vị trí (§IV) ⇒ κ = 0 theo cấu trúc; lợi thế LO đến "
                   "từ phụ thuộc mức tín hiệu",
         decomp="Không; dự báo tốt hơn (tốc độ chậm) làm LO lợi hơn hysteresis (Fig. 4–5)",
         diff="Method: tách kênh độ rộng (K2 − SC) khi κ > 0 — V–K không làm được vì phương sai hằng; phân rã; oracle "
              "một bước cùng thông tin. Problem: delay hàng đợi dưới telemetry cũ + nhiễu (đã kiểm full text)",
         no_claim="'So luật dùng xác suất với luật tĩnh đã tune, cùng dự báo, cùng ngân sách'; 'luật dựa mô hình thích "
                  "nghi môi trường tốt hơn hysteresis cố định'; tương đương ràng buộc–Bayes",
         fulltext="2026-09-30; §I–VI + Phụ lục"),
    dict(paper="Cisco SD-WAN Policies Config Guide vEdge 20.x — chương Application-Aware Routing",
         role="Nguồn thực tế SD-WAN (tài liệu nhà cung cấp)",
         vs_tuned=NA, oracle=NA,
         mechanism="Telemetry thực: BFD 1 s trên MỌI tunnel; trung bình poll 10 phút; SLA trên 6 poll (1 giờ) — để damping",
         decomp=NA,
         diff="Luật thực tế = ngưỡng SLA tuyệt đối trên trung bình dài + dải variance + tunnel dampening — đúng họ "
              "'tĩnh + damping' đề tài dùng làm baseline",
         no_claim="Không dùng 'SD-WAN đo path phụ thưa hơn' cho cấu hình mặc định",
         fulltext="2026-09-30; toàn chương (HTML)"),
    dict(paper="Richardson 2000 — Skill and relative economic value of the ECMWF EPS (QJRMS 126:649–667); Murphy 1977",
         role="Giá trị kinh tế của dự báo — cost–loss, REV",
         vs_tuned="Một phần: xác suất EPS vs dự báo tất định cùng mô hình (abstract); có tune ngưỡng tất định không — chưa kiểm",
         oracle="Chuẩn hoá theo dự báo hoàn hảo (thông tin hoàn hảo, không cùng thông tin)",
         mechanism="Không", decomp="Không (một tỉ lệ REV)",
         diff="Method: phân rã bốn bậc + baseline tâm đã tune — chưa kiểm full text Richardson",
         no_claim="'Chuẩn hoá giá trị quyết định theo thông tin hoàn hảo' (tỉ lệ headroom ~ REV)",
         fulltext="Chưa; abstract + định nghĩa REV qua ≥ 3 nguồn thứ cấp 2026-09-30"),
    dict(paper="Tài liệu Linux kernel — MPTCP sysctl stale_loss_cnt", role="Nguồn thực tế (ĐÍNH CHÍNH L1.8)",
         vs_tuned=NA, oracle=NA, mechanism=NA, decomp=NA, diff=NA,
         no_claim="Không dùng 'MPTCP stale subflow' làm bằng chứng path phụ đo thưa: 'stale' = subflow không tiến "
                  "triển, scheduler bỏ qua",
         fulltext="2026-09-30; tài liệu gốc (docs.kernel.org)"),
]


def main() -> None:
    wb = load_workbook(WB)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws.append([HEADERS[k] for k in KEYS])
    for row in ROWS:
        if set(row) != set(KEYS):
            raise ValueError(f"Thiếu/thừa cột ở dòng: {row.get('paper')}")
        ws.append([row[k] for k in KEYS])
    wrap = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for line in ws.iter_rows():
        for cell in line:
            cell.alignment = wrap
    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "B2"
    wb.save(WB)
    print(f"Đã ghi {len(ROWS)} dòng vào sheet '{SHEET}' của {WB.name}")


if __name__ == "__main__":
    main()
